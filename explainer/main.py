import asyncio
import json
import os
from pathlib import Path

from pptx import Presentation
from google import genai
from database.db_connection import SessionLocal
from database.models import User, Upload, Status


BASE_DIR = Path(__file__).resolve().parent.parent
UPLOADS_DIRECTORY = BASE_DIR / "uploads"
OUTPUTS_DIRECTORY = BASE_DIR / "outputs"

UPLOADS_DIRECTORY.mkdir(parents=True, exist_ok=True)
OUTPUTS_DIRECTORY.mkdir(parents=True, exist_ok=True)


GEMINI_API_KEY = "AQ.Ab8RN6INOnVVoOO-VSIMD3d26onNEywTdoxo2TSLO3n7nD6sCg"

if not GEMINI_API_KEY:
    raise ValueError("La variable GEMINI_API_KEY n'est pas définie.")

client = genai.Client(api_key=GEMINI_API_KEY)
# ------------------- Create a session
db=SessionLocal()


def extract_slide_text(slide):
    text_runs = []

    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()

                if text:
                    text_runs.append(text)

    return text_runs


async def get_prompt(text_runs):
    slide_text = "\n".join(text_runs)

    prompt = f"""
You are an expert presentation assistant.
Your task is to explain a PowerPoint slide clearly and simply.

Explain:
- What the slide is about
- The main idea
- The important concepts
- Any technical terms that need clarification
- What the presenter should understand from this slide

Do not invent information that is not present in the slide.

PowerPoint slide content:
-------------------------
{slide_text}
-------------------------

Return only the explanation.
"""

    try:
        response = await client.aio.models.generate_content(
            model="gemini-3.5-flash",
            contents=prompt,
        )

        if response.text:
            return response.text.strip()

        return "Gemini did not return any explanation."

    except Exception as e:
        print(f"Erreur lors de l'appel à Gemini...")
        return f"Gemini error: {str(e)}"


async def handle_gemini_explainer():

    print("Explainer start !")
    print(f"Upload directory : {UPLOADS_DIRECTORY}")
    print(f"Output directory : {OUTPUTS_DIRECTORY}")

    while True:

        upload = db.query(Upload).filter(
            Upload.status == Status.PENDING
        ).first()

        if not upload:
            print("No pending upload.")
            await asyncio.sleep(10)
            continue

        file = UPLOADS_DIRECTORY / f"{upload.uid}.pptx"

        print(f"File : {file}")

        upload.status = Status.PROCESSING
        db.commit()

        try:
            prs = Presentation(file)

            slides_data = []

            for i, slide in enumerate(prs.slides, start=1):

                text_runs = extract_slide_text(slide)

                print(f"Scan slide {i}")

                if not text_runs:
                    analysis = "No text found on this slide."
                else:
                    analysis = await get_prompt(text_runs)
                    await asyncio.sleep(14)

                slides_data.append({
                    "slide": i,
                    "text_extracted": text_runs,
                    "gemini_analysis": analysis
                })

            output_path = OUTPUTS_DIRECTORY / f"{upload.uid}.json"

            with open(output_path, "w", encoding="utf-8") as f:
                json.dump(
                    slides_data,
                    f,
                    ensure_ascii=False,
                    indent=4
                )

            upload.status = Status.DONE
            upload.finish_time = datetime.now()
            db.commit()

            print(f"Analyse enregistrée dans : {output_path}")

        except Exception as e:
            print(f"Erreur lors du traitement de {file}: {e}")

            upload.status = Status.FAILED
            db.commit()

        await asyncio.sleep(10)


if __name__ == "__main__":
    try:
        asyncio.run(handle_gemini_explainer())

    except KeyboardInterrupt:
        print("Arrêt du service.")