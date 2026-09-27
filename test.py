from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pathlib import Path

# 1. Création de la présentation
prs = Presentation()

# Définition d'un modèle de mise en page vide (Index 6 = Blank)
blank_slide_layout = prs.slide_layouts[6]

# ---------------------------------------------------------
# SLIDE 1 : Titre et Définition Générale
# ---------------------------------------------------------
slide1 = prs.slides.add_slide(blank_slide_layout)

# Titre Slide 1
tx_box1 = slide1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.5), Inches(1.2))
tf1 = tx_box1.text_frame
p1 = tf1.paragraphs[0]
p1.text = "La Photosynthèse : Définition et Rôle"
p1.font.size = Pt(32)
p1.font.bold = True
p1.font.color.rgb = RGBColor(34, 139, 34)  # Vert forêt

# Contenu Slide 1
content_box1 = slide1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(8.5), Inches(4.5))
tf_content1 = content_box1.text_frame
tf_content1.word_wrap = True

p_intro = tf_content1.paragraphs[0]
p_intro.text = "Qu'est-ce que la photosynthèse ?"
p_intro.font.size = Pt(20)
p_intro.font.bold = True

bullets_s1 = [
    "Processus biochimique essentiel réalisé par les plantes vertes, les algues et certaines bactéries.",
    "Permet de convertir l'énergie lumineuse (soleil) en énergie chimique (sucre).",
    "Prend place dans les chloroplastes des cellules végétales grâce à la chlorophylle.",
    "Essentiel pour la vie sur Terre : produit le dioxygène (O2) que nous respirons."
]

for item in bullets_s1:
    p = tf_content1.add_paragraph()
    p.text = f"• {item}"
    p.font.size = Pt(16)

# ---------------------------------------------------------
# SLIDE 2 : La Réaction Chimique et les Étapes
# ---------------------------------------------------------
slide2 = prs.slides.add_slide(blank_slide_layout)

# Titre Slide 2
tx_box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.5), Inches(1.2))
tf2 = tx_box2.text_frame
p2 = tf2.paragraphs[0]
p2.text = "Équation Chimique et Fonctionnement"
p2.font.size = Pt(32)
p2.font.bold = True
p2.font.color.rgb = RGBColor(34, 139, 34)

# Contenu Slide 2
content_box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(8.5), Inches(4.5))
tf_content2 = content_box2.text_frame
tf_content2.word_wrap = True

p_eq_title = tf_content2.paragraphs[0]
p_eq_title.text = "L'équation globale :"
p_eq_title.font.size = Pt(20)
p_eq_title.font.bold = True

p_eq = tf_content2.add_paragraph()
p_eq.text = "6 CO2 + 6 H2O + Lumière ➔ C6H12O6 (Glucose) + 6 O2"
p_eq.font.size = Pt(18)
p_eq.font.bold = True
p_eq.font.color.rgb = RGBColor(0, 102, 204)  # Bleu

p_steps_title = tf_content2.add_paragraph()
p_steps_title.text = "\nLes deux grandes phases :"
p_steps_title.font.size = Pt(20)
p_steps_title.font.bold = True

bullets_s2 = [
    "Phase lumineuse (dépendante de la lumière) : Capture l'énergie solaire et sépare l'eau (H2O), libérant ainsi du dioxygène (O2).",
    "Phase sombre (Cycle de Calvin) : Utilise l'énergie accumulée pour transformer le dioxyde de carbone (CO2) en glucose."
]

for item in bullets_s2:
    p = tf_content2.add_paragraph()
    p.text = f"• {item}"
    p.font.size = Pt(16)

# ---------------------------------------------------------
# Enregistrement du fichier
# ---------------------------------------------------------
output_path = Path("client/photosynthese_explanation.pptx")
prs.save(output_path)
print(f"Présentation enregistrée avec succès : {output_path.resolve()}")